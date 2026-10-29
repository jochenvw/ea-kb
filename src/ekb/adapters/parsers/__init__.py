from __future__ import annotations

import asyncio
import io
import json
import re
from dataclasses import dataclass, field
from importlib.metadata import version
from typing import Any

import pymupdf
from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ekb.domain.models import (
    Segment,
    SourceItem,
    SourceRevision,
    fingerprint,
    parsing_fingerprint,
    stable_id,
)
from ekb.observability import tracer


@dataclass
class Block:
    text: str
    heading: str | None = None
    level: int | None = None
    location: dict[str, Any] = field(default_factory=dict)
    quality: str = "embedded"
    media: tuple[str, ...] = ()


def split_text(text: str, max_chars: int) -> list[tuple[int, str]]:
    """Prefer paragraph/sentence boundaries; offsets always refer to the parsed text."""
    pieces: list[tuple[int, str]] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            candidates = [
                text.rfind("\n\n", start + max_chars // 3, end),
                text.rfind(". ", start + max_chars // 3, end),
                text.rfind("\n", start + max_chars // 3, end),
            ]
            boundary = next((point for point in candidates if point >= 0), -1)
            if boundary >= 0:
                end = boundary + (1 if text[boundary] == "." else 0)
        raw = text[start:end]
        trimmed = raw.strip()
        if trimmed:
            pieces.append((start + len(raw) - len(raw.lstrip()), trimmed))
        start = end
    return pieces


class StructuredParser:
    def __init__(self, *, ocr: bool = False, max_chars: int = 12_000):
        self.ocr = ocr
        self.max_chars = max_chars
        self.version = fingerprint(
            implementation="structured-v3",
            ocr=ocr,
            max_chars=max_chars,
            packages={name: version(name) for name in ("pymupdf", "python-docx", "python-pptx")},
        )

    async def parse(
        self, item: SourceItem, revision: SourceRevision, content: bytes
    ) -> list[Segment]:
        return await asyncio.to_thread(self._parse, item, revision, content)

    def _parse(self, item: SourceItem, revision: SourceRevision, content: bytes) -> list[Segment]:
        match item.format:
            case "pdf":
                blocks = self._pdf(content, item.metadata.get("pages", []))
            case "docx":
                blocks = self._docx(content)
            case "pptx":
                blocks = self._pptx(content)
            case "json" | "oscal":
                blocks = self._oscal(content, item.metadata.get("control_ids", []))
            case "web-json":
                blocks = self._web_json(content)
            case "xml":
                blocks = self._xml(content)
            case "html":
                blocks = self._html(content)
            case "txt" | "md":
                blocks = self._text(content.decode("utf-8-sig"))
            case _:
                raise ValueError(f"Unsupported source format: {item.format}")
        segments: list[Segment] = []
        parents: list[tuple[int, Any]] = []
        with tracer.start_as_current_span("segmentation"):
            for block in blocks:
                for offset, text in split_text(block.text, self.max_chars):
                    segment_id = stable_id(
                        revision.id, parsing_fingerprint(self.version, item), len(segments), text
                    )
                    if block.level is not None:
                        while parents and parents[-1][0] >= block.level:
                            parents.pop()
                    parent_id = parents[-1][1] if parents else None
                    segments.append(
                        Segment(
                            id=segment_id,
                            revision_id=revision.id,
                            ordinal=len(segments),
                            text=text,
                            heading=block.heading,
                            parent_id=parent_id,
                            location={**block.location, "block_offset": offset},
                            quality=block.quality,
                            media=block.media,
                        )
                    )
                    if block.level is not None:
                        parents.append((block.level, segment_id))
        return segments

    def _text(self, text: str) -> list[Block]:
        blocks = []
        lines: list[str] = []
        heading = None
        level = None
        start = 1
        for line_number, line in enumerate(text.splitlines(), 1):
            match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if match:
                if lines:
                    blocks.append(Block("\n".join(lines), heading, level, {"line": start}))
                lines = [line]
                heading, level, start = match[2], len(match[1]), line_number
            else:
                lines.append(line)
        if lines:
            blocks.append(Block("\n".join(lines), heading, level, {"line": start}))
        return blocks

    def _pdf(self, content: bytes, pages: list[int]) -> list[Block]:
        blocks: list[Block] = []
        heading = None
        with pymupdf.open(stream=content, filetype="pdf") as document:
            for page_index in range(len(document)):
                if pages and page_index + 1 not in pages:
                    continue
                page = document[page_index]
                textpage = None
                embedded = str(page.get_text("text", sort=True))
                quality = "embedded"
                if len(re.sub(r"\W", "", embedded)) < 20:
                    if self.ocr:
                        with tracer.start_as_current_span(
                            "ocr", attributes={"page": page_index + 1}
                        ):
                            textpage = page.get_textpage_ocr(language="eng", full=True)
                        quality = "ocr"
                    else:
                        quality = "low_text_ocr_disabled"
                media = tuple(
                    f"pdf:page:{page_index + 1}:xref:{image[0]}" for image in page.get_images()
                )
                layout: Any = page.get_text("dict", sort=True, textpage=textpage)
                found_text = False
                for block_index, block in enumerate(layout.get("blocks", [])):
                    if block.get("type") != 0:
                        continue
                    spans = [
                        span for line in block.get("lines", []) for span in line.get("spans", [])
                    ]
                    text = "\n".join(
                        "".join(span.get("text", "") for span in line.get("spans", []))
                        for line in block.get("lines", [])
                    ).strip()
                    if not text:
                        continue
                    found_text = True
                    is_heading = len(text) < 160 and any(
                        span.get("size", 0) >= 14 for span in spans
                    )
                    if is_heading:
                        heading = text
                    blocks.append(
                        Block(
                            text,
                            heading=heading,
                            level=2 if is_heading else None,
                            location={
                                "page": page_index + 1,
                                "block": block_index,
                                "bbox": list(block["bbox"]),
                                "heading_inferred": is_heading,
                            },
                            quality=quality,
                            media=media,
                        )
                    )
                if not found_text and media:
                    blocks.append(
                        Block(
                            "[Image-only page; OCR required]",
                            location={"page": page_index + 1},
                            quality="ocr_no_text" if self.ocr else "image_only",
                            media=media,
                        )
                    )
        return blocks

    def _docx(self, content: bytes) -> list[Block]:
        document = Document(io.BytesIO(content))
        blocks: list[Block] = []
        heading = None
        for index, block in enumerate(document.iter_inner_content()):
            if isinstance(block, Paragraph):
                style = block.style.name if block.style else ""
                match = re.search(r"Heading (\d+)", style or "")
                level = int(match[1]) if match else None
                if level:
                    heading = block.text
                images = tuple(
                    str(
                        node.get(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                        )
                    )
                    for node in block._p.xpath(".//a:blip")
                )
                if block.text.strip() or images:
                    blocks.append(
                        Block(
                            block.text or "[Image; see original document]",
                            heading,
                            level,
                            {"paragraph": index},
                            quality="embedded" if block.text.strip() else "image_only",
                            media=images,
                        )
                    )
            elif isinstance(block, Table):
                for row_index, row in enumerate(block.rows):
                    blocks.append(
                        Block(
                            " | ".join(cell.text for cell in row.cells),
                            heading,
                            location={"table": index, "row": row_index},
                        )
                    )
        return blocks

    def _pptx(self, content: bytes) -> list[Block]:
        presentation = Presentation(io.BytesIO(content))
        blocks = []
        for index, slide in enumerate(presentation.slides, 1):
            heading = slide.shapes.title.text if slide.shapes.title else None
            texts = []
            media = []

            def visit(
                shapes: Any,
                texts: list[str] = texts,
                media: list[str] = media,
                index: int = index,
            ) -> None:
                for shape in sorted(shapes, key=lambda shape: (shape.top, shape.left)):
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        visit(shape.shapes)
                    elif shape.has_text_frame:
                        texts.append(shape.text)
                    elif shape.has_table:
                        texts.extend(
                            " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
                        )
                    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        media.append(f"pptx:slide:{index}:shape:{shape.shape_id}")

            visit(slide.shapes)
            if texts or media:
                blocks.append(
                    Block(
                        "\n".join(texts) or "[Image-only slide]",
                        heading,
                        1,
                        {"slide": index},
                        quality="embedded" if texts else "image_only",
                        media=tuple(media),
                    )
                )
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes and notes.text.strip():
                    blocks.append(
                        Block(notes.text, heading, location={"slide": index, "notes": True})
                    )
        return blocks

    def _oscal(self, content: bytes, control_ids: list[str]) -> list[Block]:
        catalog = json.loads(content).get("catalog")
        if not isinstance(catalog, dict):
            raise ValueError("Expected an OSCAL catalog object")
        blocks = []

        def visit(node: dict[str, Any], control_id: str | None = None) -> None:
            cid = node.get("id", control_id) if "title" in node else control_id
            if node.get("prose") and (not control_ids or cid in control_ids):
                key = node.get("id") or node.get("name", "prose")
                blocks.append(
                    Block(
                        node["prose"],
                        cid,
                        location={"control_id": cid, "part_id": key, "structured_control": True},
                    )
                )
            for field_name in ("groups", "controls", "parts"):
                for child in node.get(field_name, []):
                    visit(child, cid)

        visit(catalog)
        if not blocks:
            raise ValueError("OSCAL catalog contains no control prose")
        return blocks

    def _web_json(self, content: bytes) -> list[Block]:
        data = json.loads(content)
        blocks: list[Block] = []

        def visit(value: Any, path: str = "$") -> None:
            if isinstance(value, dict):
                for key, child in value.items():
                    child_path = f"{path}.{key}"
                    if key in {
                        "text",
                        "description",
                        "title",
                        "heading",
                        "body",
                        "caption",
                    } and isinstance(child, str):
                        text = BeautifulSoup(child, "html.parser").get_text(" ", strip=True)
                        if text:
                            blocks.append(Block(text, location={"json_path": child_path}))
                    else:
                        visit(child, child_path)
            elif isinstance(value, list):
                for index, child in enumerate(value):
                    visit(child, f"{path}[{index}]")

        visit(data)
        if not blocks:
            raise ValueError("Public web JSON contains no recognized textual fields")
        return blocks

    def _xml(self, content: bytes) -> list[Block]:
        # BeautifulSoup's XML backend does not resolve external entities or access networks.
        soup = BeautifulSoup(content, "xml")
        elements = soup.find_all(["Text", "P"])
        return [
            Block(element.get_text(" ", strip=True), location={"xml_element": index})
            for index, element in enumerate(elements)
            if element.get_text(strip=True)
        ]

    def _html(self, content: bytes) -> list[Block]:
        soup = BeautifulSoup(content, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        blocks = []
        heading = None
        main = soup.find("main") or soup.find("article") or soup
        for tag in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr"]):
            text = tag.get_text(" ", strip=True)
            if not text:
                continue
            level = int(tag.name[1]) if re.fullmatch(r"h[1-4]", tag.name) else None
            if level:
                heading = text
            blocks.append(
                Block(text, heading, level, {"element": len(blocks), "html_id": tag.get("id")})
            )
        return blocks
