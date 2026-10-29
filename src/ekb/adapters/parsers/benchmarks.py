from __future__ import annotations

import io
import json
import re
import zipfile
from itertools import chain
from typing import Any
from urllib.parse import unquote

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from ekb.domain.benchmarks import (
    BenchmarkAsset,
    BenchmarkControl,
    BenchmarkEdition,
    BenchmarkPassage,
    BenchmarkSubcontrol,
    BenchmarkSuccessor,
    FrameworkReference,
    NormalizedBenchmarkDocument,
    PolicyDefinition,
    PolicyReference,
    PublishedBenchmarkMapping,
)
from ekb.domain.models import (
    PublicationEvidence,
    PublicationStatus,
    Segment,
    SourceItem,
    content_hash,
    fingerprint,
    stable_id,
)
from ekb.domain.publication import assess_publication
from ekb.observability import tracer

BENCHMARK_PARSER_VERSION = fingerprint(implementation="microsoft-benchmark-v4", workbook="3.1")
CONTROL_HEADING = re.compile(r"^## ([A-Z]{1,5}-[0-9]+):[ \t]*(.+)$", re.M)
POLICY_ID = re.compile(
    r"/providers/Microsoft\.Authorization/policyDefinitions/[0-9a-f-]{36}"
    r"(?:/versions/[0-9A-Za-z.+-]+)?",
    re.I,
)


def links(text: str) -> list[tuple[str, str]]:
    return re.findall(r"\[([^\]\n]*)\]\(([^)\s]+)\)", text)


def _key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _parts(text: str) -> dict[str, tuple[str, str]]:
    pattern = re.compile(r"^### ([^\n]+)\n|^\*\*([^*\n]+)\*\*[ \t]*:?[ \t]*", re.M)
    matches = list(pattern.finditer(text))
    result = {}
    for index, match in enumerate(matches):
        label = (match[1] or match[2]).rstrip(":").strip()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        result[_key(label)] = (label, text[match.end() : end].strip())
    return result


def _field(parts: dict[str, tuple[str, str]], key: str) -> str | None:
    return parts.get(key, ("", ""))[1] or None


def _framework(label: str) -> tuple[str, str | None] | None:
    if re.search(r"NIST.*800[\s-]*53", label, re.I):
        revision = re.search(r"(?:Rev\.?\s*|r)([45])", label, re.I)
        return "nist-sp-800-53", revision[1] if revision else None
    if re.search(r"NIST\s+CSF", label, re.I):
        version = re.search(r"v?([0-9]+\.[0-9]+)", label)
        return "nist-csf", version[1] if version else None
    if re.search(r"ISO.*27001", label, re.I):
        year = re.search(r"\b(20[0-9]{2})\b", label)
        return "iso-iec-27001", year[1] if year else None
    if re.search(r"CIS\s+Controls", label, re.I):
        version = re.search(r"v([0-9]+(?:\.[0-9]+)*)", label, re.I)
        return "cis-controls", version[1] if version else None
    if re.search(r"PCI-DSS", label, re.I):
        version = re.search(r"v([0-9]+(?:\.[0-9]+)*)", label, re.I)
        return "pci-dss", version[1] if version else None
    if re.search(r"SOC\s*2", label, re.I):
        return "soc-2", None
    return None


def _identifiers(framework: str, value: str) -> list[str]:
    patterns = {
        "nist-sp-800-53": r"(?:^|[,;\n])\s*([A-Z]{2}-[0-9]+(?:\([0-9]+\))?)",
        "nist-csf": r"(?:^|[,;\n])\s*([A-Z]{2}\.[A-Z]{2}-[0-9]{2})",
        "iso-iec-27001": r"(?:^|[,;\n])\s*(A\.[0-9]+(?:\.[0-9]+)+)",
        "cis-controls": r"(?:^|[,;\n])\s*([0-9]+(?:\.[0-9]+)+)",
        "pci-dss": r"(?:^|[,;\n])\s*([0-9]+(?:\.[0-9]+)+)",
        "soc-2": r"(?:^|[,;\n])\s*(CC[0-9]+(?:\.[0-9]+)+)",
    }
    return list(dict.fromkeys(re.findall(patterns[framework], value)))


def _workbook_text(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 200_000_000:
            raise ValueError("Workbook expands beyond the validation limit")
        if "[Content_Types].xml" not in archive.namelist():
            raise ValueError("Expected a valid workbook package")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    chunks = []
    try:
        for sheet in book.worksheets:
            if (sheet.max_row or 0) > 20_000 or (sheet.max_column or 0) > 100:
                raise ValueError("Workbook dimensions exceed the validation limit")
            rows = sheet.iter_rows(values_only=True)
            first = next(rows, ())
            headers = [str(value).strip() if value is not None else "" for value in first]
            keys = [_key(value) for value in headers]
            identifier = next(
                (i for i, key in enumerate(keys) if key in {"id", "azureid", "controlid"}), None
            )
            if identifier is None:
                if sheet.title.casefold() == "readme":
                    for row in chain((first,), rows):
                        chunks.extend(str(value).strip() + "\n" for value in row if value)
                continue
            title_index = next(
                (
                    i
                    for i, key in enumerate(keys)
                    if key
                    in {"recommendation", "benchmarkrecommendation", "controlname", "controltitle"}
                ),
                None,
            )
            if title_index is None:
                continue
            for row_number, row in enumerate(rows, 2):
                cid = str(row[identifier] or "").strip()
                if not re.fullmatch(r"[A-Z]{1,5}-[0-9]+", cid):
                    continue
                title = str(row[title_index] or "").strip()
                chunks.append(f"## {cid}: {title}\n\nWorksheet: {sheet.title}; row: {row_number}\n")
                for header, value in zip(headers, row, strict=False):
                    if header and value is not None and str(value).strip():
                        chunks.append(f"### {header}\n\n{str(value).strip()}\n")
        if not any(CONTROL_HEADING.search(chunk) for chunk in chunks):
            raise ValueError("Workbook contains no recognizable benchmark controls")
        return "\n".join(chunks)
    finally:
        book.close()


def _html_text(content: bytes) -> str:
    soup = BeautifulSoup(content, "html.parser")
    main = soup.find("main") or soup.find("article") or soup
    for tag in main.find_all(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    chunks = []
    for tag in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr"]):
        fragment = BeautifulSoup(str(tag), "html.parser")
        for anchor in fragment.find_all("a", href=True):
            anchor.replace_with(f"[{anchor.get_text(' ', strip=True)}]({anchor.get('href')})")
        for bold in fragment.find_all(["strong", "b"]):
            bold.replace_with(f"**{bold.get_text(' ', strip=True)}**")
        if tag.name == "tr":
            cells = fragment.find_all(["td", "th"])
            text = "| " + " | ".join(cell.get_text(" ", strip=True) for cell in cells) + " |"
            chunks.append(text)
            if tag.find("th"):
                chunks.append("| " + " | ".join("---" for _ in cells) + " |")
        else:
            text = fragment.get_text(" ", strip=True)
            prefix = "#" * int(tag.name[1]) + " " if re.fullmatch(r"h[1-4]", tag.name) else ""
            chunks.append(prefix + text)
    return "\n".join(chunks)


def _policy(content: bytes, asset: BenchmarkAsset) -> PolicyDefinition:
    data = json.loads(content)
    properties = data.get("properties", {})
    definition_id = data.get("id", "")
    if not POLICY_ID.fullmatch(definition_id) or not properties.get("displayName"):
        raise ValueError("Expected an actual Azure Policy definition")
    types: list[str] = []

    def visit(value: Any) -> None:
        if isinstance(value, dict):
            if value.get("field") == "type":
                if isinstance(value.get("equals"), str):
                    types.append(value["equals"])
                if isinstance(value.get("in"), list):
                    types.extend(item for item in value["in"] if isinstance(item, str))
            for child in value.values():
                visit(child)
        elif isinstance(value, list):
            for child in value:
                visit(child)

    visit(properties.get("policyRule", {}).get("if", {}))
    return PolicyDefinition(
        definition_id=definition_id,
        name=properties["displayName"],
        version=properties.get("metadata", {}).get("version"),
        resource_types=tuple(dict.fromkeys(types)),
        mode=properties.get("mode"),
        metadata=properties.get("metadata", {}),
        parameters=properties.get("parameters", {}),
        rule=properties.get("policyRule", {}),
        source_url=asset.url,
        content_hash=content_hash(content),
    )


def normalize_benchmark(
    asset: BenchmarkAsset, content: bytes, edition: BenchmarkEdition
) -> NormalizedBenchmarkDocument:
    with tracer.start_as_current_span("benchmark.parse"):
        if asset.format == "xlsx":
            text = _workbook_text(content)
        elif asset.format == "html":
            text = _html_text(content)
        else:
            text = content.decode("utf-8-sig").replace("\r\n", "\n")
    digest = content_hash(content)
    segment_id = stable_id(asset.id, digest, BENCHMARK_PARSER_VERSION, "benchmark-text")

    def passage(excerpt: str, section: str, start: int = 0) -> BenchmarkPassage:
        position = text.index(excerpt, start)
        return BenchmarkPassage(
            source_id=asset.id,
            source_url=asset.url,
            content_hash=digest,
            section=section,
            excerpt=excerpt,
            start=position,
            end=position + len(excerpt),
        )

    title_match = re.search(r"^# (.+)$", text, re.M)
    title = title_match[1] if title_match else asset.url.rsplit("/", 1)[-1]
    policies = (_policy(content, asset),) if asset.role == "policy" else ()
    if policies:
        title = policies[0].name
    publisher_version = (
        policies[0].version
        if policies
        else None
        if asset.role in {"policy_index", "repository_index"}
        else edition.publisher_version
    )
    publication = assess_publication(
        SourceItem(
            id=asset.id,
            connector="microsoft-security-benchmark",
            location=asset.url,
            name=title,
            format=asset.format,
            metadata={"publisher_version": publisher_version},
        ),
        [Segment(id=segment_id, revision_id=stable_id(asset.id, digest), ordinal=0, text=text)],
    )
    workbook_preview = re.search(
        r"^This spreadsheet is designed to provide[^\n]*\bprivate preview\b[^\n]*", text, re.M
    )
    if workbook_preview:
        publication = publication.model_copy(
            update={
                "publication_status": PublicationStatus.PREVIEW,
                "evidence": (
                    *publication.evidence,
                    PublicationEvidence(
                        excerpt=workbook_preview[0],
                        segment_id=segment_id,
                        start=workbook_preview.start(),
                        end=workbook_preview.end(),
                    ),
                ),
            }
        )

    def mappings(block: str, cid: str, offset: int) -> tuple[FrameworkReference, ...]:
        found = []

        def append(label: str, value: str, excerpt: str) -> None:
            framework = _framework(label)
            if framework is None:
                return
            for identifier in _identifiers(framework[0], value):
                found.append(
                    FrameworkReference(
                        framework=framework[0],
                        revision=framework[1],
                        control_id=identifier,
                        publisher_label=label,
                        evidence=passage(excerpt, cid, offset),
                    )
                )

        with tracer.start_as_current_span("benchmark.extract_mappings"):
            for line in block.splitlines():
                match = re.match(r"- \*\*([^*]+):\*\*\s*(.+)", line)
                if match:
                    append(match[1], match[2], line)
            lines = block.splitlines()
            for i in range(len(lines) - 2):
                if lines[i].startswith("|") and re.fullmatch(r"[| :\-]+", lines[i + 1]):
                    headers = lines[i].strip("|").split("|")
                    values = lines[i + 2].strip("|").split("|")
                    for label, value in zip(headers, values, strict=False):
                        append(label.strip(), value.strip(), "\n".join(lines[i : i + 3]))
            for label, value in _parts(block).values():
                if _framework(label):
                    excerpt = f"### {label}\n\n{value}"
                    if excerpt in text:
                        append(label, value, excerpt)
        return tuple({(ref.framework, ref.revision, ref.control_id): ref for ref in found}.values())

    controls = []
    with tracer.start_as_current_span("benchmark.extract_controls"):
        if asset.role in {"domain", "workbook"}:
            for match in CONTROL_HEADING.finditer(text):
                following = re.search(r"^## ", text[match.end() :], re.M)
                end = match.end() + following.start() if following else len(text)
                block = text[match.start() : end].strip()
                cid, name = match[1], match[2].strip()
                parts = _parts(block)

                principle = _field(parts, "securityprinciple")
                subcontrols = tuple(
                    BenchmarkSubcontrol(
                        control_id=label.split(":", 1)[0].strip(),
                        title=label.split(":", 1)[1].strip(),
                        text=value,
                        evidence=passage(f"### {label}\n\n{value}", label, match.start()),
                    )
                    for label, value in parts.values()
                    if re.match(rf"^{re.escape(cid)}\.[0-9]+:", label)
                )
                guidance = _field(parts, "azureguidance") or "\n\n".join(
                    subcontrol.text for subcontrol in subcontrols
                )
                for additional in (
                    "azureimplementationandadditionalcontext",
                    "implementationexample",
                ):
                    if value := _field(parts, additional):
                        guidance += "\n\n" + value
                options = tuple(
                    option
                    for option in (
                        "Azure Application Gateway",
                        "Azure Front Door",
                        "Azure Content Delivery Network",
                        "Azure CDN",
                    )
                    if option in guidance
                    and not (option == "Azure CDN" and "Azure Content Delivery Network" in guidance)
                )
                alternatives = any(
                    sum(option in paragraph for option in options) > 1
                    and re.search(r"\b(?:or|choose|select)\b", paragraph, re.I)
                    for paragraph in re.split(r"\n\n|(?<=\.)\s+", guidance)
                )
                controls.append(
                    BenchmarkControl(
                        benchmark=edition.family,
                        edition=edition.key,
                        publisher_version=edition.publisher_version,
                        domain=cid.split("-")[0],
                        control_id=cid,
                        title=name,
                        security_principle=principle,
                        generic_guidance=_field(parts, "genericguidance")
                        or _field(parts, "guidance"),
                        risk_to_mitigate=_field(parts, "risktomitigate"),
                        azure_guidance=guidance,
                        criticality=_field(parts, "criticalitylevel"),
                        subcontrols=subcontrols,
                        mappings=mappings(block, cid, match.start()),
                        implementation_options=options,
                        implementation_options_are_alternatives=True if alternatives else None,
                        evidence=passage(block, cid, match.start()),
                        principle_evidence=passage(principle, cid, match.start())
                        if principle
                        else None,
                        publication=publication,
                    )
                )
            if not controls:
                raise ValueError("Expected actual benchmark control sections")

    successors = []
    if asset.role == "overview":
        for line in text.splitlines():
            if re.search(r"\b(?:supersedes|replaces)\b", line, re.I):
                for _label, target in links(line):
                    predecessor = (
                        "mcsb-v1"
                        if target.rstrip("/").endswith("overview-mcsb-v1")
                        else "asb-v2"
                        if target.rstrip("/").endswith("overview-v2")
                        else None
                    )
                    if predecessor and predecessor != edition.key:
                        successors.append(
                            BenchmarkSuccessor(
                                successor=edition.key,
                                predecessor=predecessor,
                                evidence=passage(line, "Published edition succession"),
                            )
                        )
    references = []
    section = "Policy references"
    for line in text.splitlines():
        if heading := CONTROL_HEADING.match(line):
            section = heading[1]
        for label, target in links(line):
            identifier = POLICY_ID.search(unquote(target))
            if identifier:
                cells = line.split("|")
                version = re.search(
                    r"\b[0-9]+\.[0-9]+\.[0-9]+\b", cells[-2] if len(cells) >= 3 else ""
                )
                references.append(
                    PolicyReference(
                        definition_id=identifier[0],
                        name=label,
                        url=target,
                        reference_version=version[0] if version else None,
                        evidence=passage(line, section),
                    )
                )
    published_mappings = []
    if asset.role == "mapping" and asset.format != "json":
        headers: list[str] = []
        for line in text.splitlines():
            if not line.startswith("|"):
                continue
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if any(_framework(cell) for cell in cells):
                headers = cells
                continue
            if not headers or not re.search(r"\b[A-Z]{1,5}-[0-9]+\b", cells[0]):
                continue
            for label, value in zip(headers, cells, strict=False):
                framework = _framework(label)
                if framework is None:
                    continue
                for cid in re.findall(r"\b[A-Z]{1,5}-[0-9]+\b", cells[0]):
                    for identifier in _identifiers(framework[0], value):
                        published_mappings.append(
                            PublishedBenchmarkMapping(
                                edition=edition.key,
                                control_id=cid,
                                reference=FrameworkReference(
                                    framework=framework[0],
                                    revision=framework[1],
                                    control_id=identifier,
                                    publisher_label=label,
                                    evidence=passage(line, cid),
                                ),
                            )
                        )
    return NormalizedBenchmarkDocument(
        asset_id=asset.id,
        segment_id=segment_id,
        title=title,
        content_hash=digest,
        parser_version=BENCHMARK_PARSER_VERSION,
        text=text,
        controls=tuple(controls),
        successors=tuple(successors),
        policies=policies,
        policy_references=tuple(references),
        mappings=tuple(published_mappings),
        publication=publication,
    )
