import asyncio
import json
import re
from html import escape
from pathlib import Path

import httpx
import pytest
from sqlalchemy import func, select

from ekb.adapters.artifacts import FilesystemArtifactStore
from ekb.adapters.inference import RulesInference
from ekb.adapters.parsers import StructuredParser
from ekb.adapters.persistence.postgres.models import OccurrenceRow, RevisionRow, SegmentRow
from ekb.adapters.sources import FilesystemConnector
from ekb.domain.models import (
    Candidate,
    Capabilities,
    ContentLabel,
    ContentLabeling,
    Extraction,
    GoalLevel,
    InferenceResult,
    LabelingResult,
    Modality,
    SemanticType,
    Usage,
)
from ekb.pipeline import IngestionPipeline

pytestmark = pytest.mark.integration


class CountingParser(StructuredParser):
    calls = 0

    async def parse(self, item, revision, content):
        self.calls += 1
        return await super().parse(item, revision, content)


class CountingInference(RulesInference):
    calls = 0
    fail_at: int | None = None

    async def extract(self, segment):
        self.calls += 1
        if self.calls == self.fail_at:
            raise RuntimeError("Simulated inference interruption")
        await asyncio.sleep(0.01)
        return await super().extract(segment)


class SynthesizingInference(RulesInference):
    capabilities = Capabilities(model_id="synthetic-live-model")
    strategy_reconciliation_fingerprint = "synthetic-strategy-reconciliation-v1"

    async def label(self, segment):
        labels = [
            ContentLabel(
                excerpt=text,
                semantic_type=(
                    SemanticType.STRATEGIC_REQUIREMENT
                    if "confidential information" in text.lower()
                    else SemanticType.ARCHITECTURE_DOMAIN
                    if "architecture domain" in text.lower()
                    else SemanticType.STRATEGIC_GOAL
                ),
            )
            for text in (
                "We aim to become a resilient enterprise.",
                "We aim to build trusted relationships.",
                "Grow recurring revenue by 5% by 2030.",
                "Increase recurring revenue 5% by 2030.",
                "Teams must protect confidential information.",
                "Confidential information must be protected by teams.",
                "Architecture domain: Identity management.",
            )
        ]
        return LabelingResult(
            labeling=ContentLabeling(labels=labels),
            usage=Usage(model_calls=1),
            model_id=self.capabilities.model_id,
            prompt_version="synthetic-labeling-v1",
        )

    async def extract_scoped(self, segment, labeling):
        return InferenceResult(
            extraction=Extraction(
                candidates=[
                    Candidate(
                        statement=label.excerpt,
                        excerpt=label.excerpt,
                        semantic_type=label.semantic_type,
                        goal_level=(
                            GoalLevel.SUBORDINATE_TARGET
                            if "5%" in label.excerpt
                            else GoalLevel.ENTERPRISE
                            if label.semantic_type is SemanticType.STRATEGIC_GOAL
                            else None
                        ),
                        modality=(
                            Modality.MANDATED
                            if label.semantic_type is SemanticType.STRATEGIC_REQUIREMENT
                            else Modality.UNSPECIFIED
                        ),
                    )
                    for label in labeling.labels
                ]
            ),
            usage=Usage(model_calls=1),
            model_id=self.capabilities.model_id,
            prompt_version="synthetic-extraction-v1",
        )

    async def reconcile_strategy(self, segment):
        source_key = re.search(r"Source key: (segment:[0-9a-f-]+)", segment.text)
        evidence = "We aim to become a resilient enterprise."
        target = "Grow recurring revenue by 5% by 2030."
        requirement = "Teams must protect confidential information."
        assert (
            source_key
            and evidence in segment.text
            and target in segment.text
            and requirement in segment.text
        )
        return InferenceResult(
            extraction=Extraction(
                candidates=[
                    Candidate(
                        statement="Become a resilient enterprise.",
                        excerpt=evidence,
                        semantic_type=SemanticType.STRATEGIC_GOAL,
                        goal_level=GoalLevel.ENTERPRISE,
                        source_key=source_key.group(1),
                    ),
                    Candidate(
                        statement=target,
                        excerpt=target,
                        semantic_type=SemanticType.STRATEGIC_GOAL,
                        goal_level=GoalLevel.SUBORDINATE_TARGET,
                        source_key=source_key.group(1),
                    ),
                    Candidate(
                        statement=requirement,
                        excerpt=requirement,
                        semantic_type=SemanticType.STRATEGIC_REQUIREMENT,
                        modality=Modality.MANDATED,
                        source_key=source_key.group(1),
                    ),
                ]
            ),
            usage=Usage(model_calls=1),
            model_id=self.capabilities.model_id,
            prompt_version=self.strategy_reconciliation_fingerprint,
        )


def build(database, tmp_path):
    source = tmp_path / "source"
    source.mkdir(exist_ok=True)
    parser, inference = CountingParser(), CountingInference()
    pipeline = IngestionPipeline(
        database, FilesystemArtifactStore(tmp_path / "artifacts"), parser, inference
    )
    return pipeline, parser, inference, source


async def test_live_ingestion_reconciles_strategy_and_persists_domain(database, tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "strategy.md").write_text(
        "# Strategy\n"
        "We aim to become a resilient enterprise. "
        "We aim to build trusted relationships. "
        "Grow recurring revenue by 5% by 2030. "
        "Increase recurring revenue 5% by 2030. "
        "Teams must protect confidential information. "
        "Confidential information must be protected by teams. "
        "Architecture domain: Identity management."
    )
    pipeline = IngestionPipeline(
        database,
        FilesystemArtifactStore(tmp_path / "artifacts"),
        StructuredParser(),
        SynthesizingInference(),
    )

    report = await pipeline.scan([FilesystemConnector(source)])

    assert report.processed == 1 and report.failed == 0 and report.concepts == 4
    async with database() as uow:
        concepts = await uow.knowledge.search()
    assert {
        (
            item.versions[0].candidate.statement,
            item.versions[0].candidate.goal_level,
        )
        for item in concepts
    } == {
        ("Become a resilient enterprise.", GoalLevel.ENTERPRISE),
        ("Grow recurring revenue by 5% by 2030.", GoalLevel.SUBORDINATE_TARGET),
        ("Teams must protect confidential information.", None),
        ("Architecture domain: Identity management.", None),
    }


async def test_incremental_revisions_grounding_and_deletion(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    path = source / "policy.md"
    path.write_text(
        "# Security\nOperators must retain logs.\n\n# Strategy\nOur vision is to reduce costs."
    )
    first = await pipeline.scan([FilesystemConnector(source)])
    assert first.processed == 1 and first.concepts == 2 and first.failed == 0
    calls = inference.calls
    second = await pipeline.scan([FilesystemConnector(source)])
    assert second.skipped == 1 and parser.calls == 1 and inference.calls == calls
    path.write_text(
        "# Security\nOperators must retain logs.\n\n"
        "# Strategy\nOur vision is to improve availability."
    )
    third = await pipeline.scan([FilesystemConnector(source)])
    assert third.processed == 1 and third.failed == 0
    assert parser.calls == 2 and inference.calls == calls + 1
    async with database() as uow:
        assert await uow.session.scalar(select(func.count()).select_from(RevisionRow)) == 2
        concepts = await uow.knowledge.search()
        assert len(concepts) == 3
        log_concept = next(
            c for c in concepts if "retain logs" in c.versions[0].candidate.statement
        )
        provenance = await uow.knowledge.provenance(log_concept.id)
        assert len(provenance) == 2
        for item in provenance:
            occ = item.occurrence
            assert item.segment.text[occ.start : occ.end] == occ.excerpt
            assert item.revision.source_id == item.source.item.id
    path.unlink()
    deleted = await pipeline.scan([FilesystemConnector(source)])
    assert deleted.missing == 1
    async with database() as uow:
        assert (await uow.sources.list())[0].status == "missing"
        assert len(await uow.knowledge.search()) == 3
        assert await uow.session.scalar(select(func.count()).select_from(OccurrenceRow)) == 4


async def test_ordinary_http_sources_use_shared_ontology_and_replay_without_reprocessing(
    database, tmp_path
):
    from ekb.adapters.sources import HTTPConnector, read_manifest
    from ekb.domain.models import (
        Capabilities,
        ContentLabeling,
        Extraction,
        InferenceResult,
        LabelingResult,
        content_hash,
    )
    from ekb.evaluation.skills import SkillDataset

    raw_dataset = await asyncio.to_thread(Path("evals/skills-v1.json").read_bytes)
    dataset = SkillDataset.model_validate_json(raw_dataset)
    wanted = {
        "source-network-firewall-technical",
        "source-network-firewall-platform",
        "source-nist-data-protection-conceptual",
        "test-unknown-layer",
    }
    examples = {case.id: case for case in dataset.cases if case.id in wanted}
    assert set(examples) == wanted
    real = next(s for s in read_manifest(Path("sources/public.yaml")) if s.id == "mcsb-v1-network")
    nist = next(s for s in read_manifest(Path("sources/public.yaml")) if s.id == "nist-csf-2.0")
    nist = nist.model_copy(
        update={
            "control_ids": ["PR.DS-01"],
            "enterprise_requirement": None,
            "architecture_principles": [],
        }
    )
    synthetic = real.model_copy(
        update={
            "id": "synthetic-other-publisher",
            "publisher": "Example publisher",
            "name": "Synthetic architecture guidance",
            "location": "https://example.test/architecture",
            "publisher_version": None,
            "access_notes": "Synthetic test source",
            "platform_scope": None,
        }
    )

    def page(ids):
        body = "".join(f"<p>{escape(examples[case_id].text)}</p>" for case_id in ids)
        return f"<html><body><main><h1>Guidance</h1>{body}</main></body></html>".encode()

    bodies = {
        real.location: page(
            ["source-network-firewall-technical", "source-network-firewall-platform"]
        ),
        nist.location: json.dumps(
            {
                "catalog": {
                    "groups": [
                        {
                            "id": "PR",
                            "title": "Protect",
                            "controls": [
                                {
                                    "id": "PR.DS-01",
                                    "title": "Data protection",
                                    "parts": [
                                        {
                                            "id": "PR.DS-01_statement",
                                            "prose": examples[
                                                "source-nist-data-protection-conceptual"
                                            ].text,
                                        }
                                    ],
                                }
                            ],
                        }
                    ],
                }
            }
        ).encode(),
        synthetic.location: page(["test-unknown-layer"]),
    }
    requests = []

    def handle(request):
        requests.append(request)
        if request.headers.get("if-none-match") == '"test-revision"':
            return httpx.Response(304)
        return httpx.Response(
            200, content=bodies[str(request.url)], headers={"etag": '"test-revision"'}
        )

    class FixtureSkills(RulesInference):
        capabilities = Capabilities(model_id="deterministic-fixture")
        prompt_version = "source-regression-v1"
        label_fingerprint = "source-regression-label-v1"
        extraction_fingerprint = "source-regression-extraction-v1"
        label_calls = 0
        extraction_calls = 0

        async def label(self, segment):
            self.label_calls += 1
            labels = [
                label
                for case in examples.values()
                for label in case.labeling.labels
                if label.excerpt in segment.text
            ]
            return LabelingResult(
                labeling=ContentLabeling(labels=labels),
                model_id=self.capabilities.model_id,
                prompt_version=self.prompt_version,
            )

        async def extract_scoped(self, segment, labeling):
            self.extraction_calls += 1
            candidates = [
                candidate
                for case in examples.values()
                for candidate in case.expected.candidates
                if candidate.excerpt in segment.text
            ]
            return InferenceResult(
                extraction=Extraction(candidates=candidates),
                model_id=self.capabilities.model_id,
                prompt_version=self.prompt_version,
            )

    parser, skills = CountingParser(), FixtureSkills()
    artifacts = FilesystemArtifactStore(tmp_path / "artifacts")
    pipeline = IngestionPipeline(database, artifacts, parser, skills)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handle)) as client:
        connector = HTTPConnector([real, nist, synthetic], client=client)
        first = await pipeline.scan([connector])
        assert first.failed == 0 and first.processed == 3 and first.concepts == 4
        first_calls = (parser.calls, skills.label_calls, skills.extraction_calls)
        async with database() as uow:
            sources = await uow.sources.list()
            concepts = await uow.knowledge.search(limit=50)
            before = {
                c.id: (tuple(v.id for v in c.versions), tuple(o.id for o in c.occurrences))
                for c in concepts
            }
            assert len(sources) == 3 and len(concepts) == 4
            by_excerpt = {c.versions[-1].candidate.excerpt: c for c in concepts}
            for case in examples.values():
                candidate = case.expected.candidates[0]
                concept = by_excerpt[candidate.excerpt]
                stored = concept.versions[-1].candidate
                if case.id == "source-network-firewall-technical":
                    assert stored.semantic_type == SemanticType.PLATFORM_REQUIREMENT
                    assert stored.scope.platform == "Azure, AWS and GCP"
                else:
                    assert stored.semantic_type == candidate.semantic_type
                    assert stored.scope == candidate.scope
                assert not concept.relationships and not concept.authority_events
                assert all(
                    state.governance_status == "proposed"
                    and state.canonical_status == "not_canonical"
                    for state in concept.authority_states
                )
                provenance = await uow.knowledge.provenance(concept.id)
                assert len(provenance) == 1
                p = provenance[0]
                assert p.segment.text[p.occurrence.start : p.occurrence.end] == candidate.excerpt
                assert p.revision.content_hash == content_hash(bodies[p.source.item.location])
                assert (
                    await artifacts.get(p.revision.artifact_key) == bodies[p.source.item.location]
                )
                source_view = await uow.inspection.source(p.source.item.id)
                assert source_view and any(
                    item.version.id == p.occurrence.concept_version_id for item in source_view.items
                )
                graph = await uow.inspection.graph(concept.id)
                assert graph and any(
                    node.id == concept.id and node.kind == (stored.semantic_type or "unclassified")
                    for node in graph.nodes
                )
                assert any(node.id == p.source.item.id for node in graph.nodes)
                if case.source and case.source.id == "ns1":
                    assert p.source.item.metadata["publisher"] == "Microsoft"
                    assert "CC-BY 4.0" in p.source.item.metadata["access_notes"]
                elif case.source:
                    assert p.source.item.metadata["publisher"] == (
                        "US National Institute of Standards and Technology"
                    )
                else:
                    assert p.source.item.metadata["publisher"] == "Example publisher"
            snapshot = await uow.resolution.snapshot()
            assert not snapshot.findings and not snapshot.decisions
        second = await pipeline.scan([connector])
        assert second.failed == 0 and second.skipped == 3
        assert (parser.calls, skills.label_calls, skills.extraction_calls) == first_calls
        assert len(requests) == 6 and all(
            request.headers.get("if-none-match") == '"test-revision"' for request in requests[3:]
        )
        async with database() as uow:
            concepts = await uow.knowledge.search(limit=50)
            after = {
                c.id: (tuple(v.id for v in c.versions), tuple(o.id for o in c.occurrences))
                for c in concepts
            }
            assert after == before


async def test_generic_mappings_retain_evidence_without_resolving_unknown_references(
    database, tmp_path
):
    from ekb.domain.grounding import GroundingError
    from ekb.domain.models import (
        Candidate,
        ContentLabel,
        ContentLabeling,
        Extraction,
        InferenceResult,
        LabelingResult,
        ProposedRelationship,
        RelationshipKind,
        SemanticType,
    )
    from ekb.domain.skills import validate_scoped_extraction

    mapping = "Guidance A maps 'Protect network traffic.' to 'Inspect network traffic.'"
    refinement = "The latter refines the former."
    text = (
        f"Protect network traffic. Inspect network traffic. {mapping} {refinement} "
        "Reference Z-9 is unresolved."
    )
    candidates = [
        Candidate(
            semantic_type=SemanticType.CONCEPTUAL_REQUIREMENT,
            statement="Protect network traffic.",
            excerpt="Protect network traffic.",
        ),
        Candidate(
            semantic_type=SemanticType.TECHNICAL_REQUIREMENT,
            statement="Inspect network traffic.",
            excerpt="Inspect network traffic.",
        ),
    ]
    relationships = [
        ProposedRelationship(
            from_index=0, to_index=1, kind=RelationshipKind.MAPS_TO, excerpt=mapping
        ),
        ProposedRelationship(
            from_index=1, to_index=0, kind=RelationshipKind.REFINES, excerpt=refinement
        ),
    ]

    class MappingSkills(RulesInference):
        async def label(self, segment):
            return LabelingResult(
                labeling=ContentLabeling(
                    labels=[
                        ContentLabel(semantic_type=c.semantic_type, excerpt=segment.text)
                        for c in candidates
                    ]
                ),
                model_id="mapping-fixture",
                prompt_version="mapping-v1",
            )

        async def extract_scoped(self, segment, labeling):
            return InferenceResult(
                extraction=Extraction(candidates=candidates, relationships=relationships),
                model_id="mapping-fixture",
                prompt_version="mapping-v1",
            )

    source = tmp_path / "source"
    source.mkdir()
    (source / "mapping.txt").write_text(text, encoding="utf-8")
    pipeline = IngestionPipeline(
        database,
        FilesystemArtifactStore(tmp_path / "artifacts"),
        StructuredParser(),
        MappingSkills(),
    )
    report = await pipeline.scan([FilesystemConnector(source)])
    assert report.processed == 1 and report.failed == 0

    async with database() as uow:
        concepts = await uow.knowledge.search(limit=10)
        assert len(concepts) == 2
        assert all(not concept.authority_events for concept in concepts)
        root = next(
            c for c in concepts if c.versions[0].candidate.statement == candidates[0].statement
        )
        graph = await uow.inspection.graph(root.id)
        assert graph is not None
        edges = {edge.kind: edge for edge in graph.edges if edge.kind in {"mapsTo", "refines"}}
        assert set(edges) == {"mapsTo", "refines"}
        for kind, expected in (("mapsTo", mapping), ("refines", refinement)):
            edge = edges[kind]
            assert edge.status == "proposed" and edge.evidence is not None
            assert edge.evidence.excerpt == expected
        provenance = await uow.knowledge.provenance(root.id)
        segment = provenance[0].segment
        assert not (await uow.resolution.snapshot()).decisions

    unresolved = ProposedRelationship(
        from_index=0, to_index=2, kind=RelationshipKind.MAPS_TO, excerpt=mapping
    )
    labels = ContentLabeling(
        labels=[
            ContentLabel(semantic_type=c.semantic_type, excerpt=segment.text) for c in candidates
        ]
    )
    with pytest.raises(GroundingError, match="nonexistent candidate"):
        validate_scoped_extraction(
            Extraction(candidates=candidates, relationships=[unresolved]), segment, labels
        )


async def test_failed_task_reuses_completed_segment_and_parse_checkpoint(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    (source / "policy.md").write_text(
        "# A\nUsers must authenticate.\n# B\nAdmins should review logs."
    )
    inference.fail_at = 2
    first = await pipeline.scan([FilesystemConnector(source)])
    assert first.failed == 1
    assert parser.calls == 1 and inference.calls == 2
    second = await pipeline.scan([FilesystemConnector(source)])
    assert second.processed == 1 and second.failed == 0
    assert parser.calls == 1 and inference.calls == 3
    async with database() as uow:
        tasks = await uow.processing.tasks()
        assert tasks[0].status == "done" and tasks[0].attempts == 2
        assert len(await uow.knowledge.search()) == 2


async def test_concurrent_workers_claim_one_task(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    reports = await asyncio.gather(
        pipeline.scan([FilesystemConnector(source)]),
        pipeline.scan([FilesystemConnector(source)]),
    )
    assert sum(r.failed for r in reports) == 0
    assert sum(r.processed for r in reports) == 1
    assert inference.calls == 1
    async with database() as uow:
        assert await uow.session.scalar(select(func.count()).select_from(RevisionRow)) == 1
        assert await uow.session.scalar(select(func.count()).select_from(SegmentRow)) == 1
        assert len(await uow.knowledge.search()) == 1


async def test_orchestrated_scan_accepts_task_completed_during_claim_race(
    database, tmp_path, monkeypatch
):
    pipeline, _, inference, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    parsing, release, waiting, finished = (asyncio.Event() for _ in range(4))

    class PausedParser(CountingParser):
        async def parse(self, item, revision, content):
            parsing.set()
            await release.wait()
            return await super().parse(item, revision, content)

    parser = PausedParser()
    pipeline.parser = parser
    strict = IngestionPipeline(
        database, pipeline.artifacts, parser, inference, require_complete=True
    )
    ingest = strict._ingest

    async def wait_for_completion(item, revision, content):
        waiting.set()
        await finished.wait()
        return await ingest(item, revision, content)

    monkeypatch.setattr(strict, "_ingest", wait_for_completion)
    first = asyncio.create_task(pipeline.scan([FilesystemConnector(source)]))
    second = None
    try:
        await asyncio.wait_for(parsing.wait(), 10)
        second = asyncio.create_task(strict.scan([FilesystemConnector(source)]))
        await asyncio.wait_for(waiting.wait(), 10)
        release.set()
        assert (await asyncio.wait_for(first, 30)).processed == 1
        finished.set()
        result = await asyncio.wait_for(second, 30)
        assert result.status == "completed" and result.skipped == 1
        assert parser.calls == inference.calls == 1
    finally:
        release.set()
        finished.set()
        await asyncio.gather(*(task for task in (first, second) if task), return_exceptions=True)


async def test_dry_run_has_no_writes_or_model_calls(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    report = await pipeline.scan([FilesystemConnector(source)], dry_run=True)
    assert report.processed == 1 and report.dry_run
    assert parser.calls == 0 and inference.calls == 0
    async with database() as uow:
        assert await uow.sources.list() == []
        assert await uow.processing.scans() == []


async def test_staged_acquisition_barrier_and_bounded_parallel_extraction(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    for index in range(4):
        (source / f"policy-{index}.txt").write_text(f"Users must retain audit {index}.")
    fetched = set()
    active, peak = 0, 0

    class ParallelFiles(FilesystemConnector):
        async def fetch(self, item, previous):
            nonlocal active, peak
            active += 1
            peak = max(peak, active)
            await asyncio.sleep(0.03)
            result = await super().fetch(item, previous)
            fetched.add(item.id)
            active -= 1
            return result

    original = parser.parse

    async def parse_after_acquisition(item, revision, content):
        assert len(fetched) == 4, "No document may parse before acquisition completes"
        return await original(item, revision, content)

    parser.parse = parse_after_acquisition
    phases = []
    first = await pipeline.scan(
        [ParallelFiles(source)], staged=True, phase=lambda phase, report: phases.append(phase)
    )
    assert phases == ["discovery", "discovered", "completed"]
    assert peak > 1 and peak <= pipeline.concurrency
    assert first.processed == 4 and first.unknown == 4 and first.known == 0
    calls = inference.calls
    second = await pipeline.scan([ParallelFiles(source)], staged=True)
    assert second.known == second.reused == second.skipped == 4
    assert second.unknown == second.processed == 0 and inference.calls == calls


async def test_staged_http_downloads_count_real_transfers_and_reuse_304(database, tmp_path):
    import httpx

    from ekb.adapters.sources import HTTPConnector
    from ekb.domain.models import SourceSpec

    pipeline, _, inference, _ = build(database, tmp_path)
    requests = []

    def respond(request):
        requests.append(request)
        if request.headers.get("if-none-match") == '"stable"':
            return httpx.Response(304)
        return httpx.Response(200, text="Users must authenticate.", headers={"etag": '"stable"'})

    specification = SourceSpec(
        id="synthetic-download",
        publisher="Example",
        name="Synthetic policy",
        location="https://example.invalid/policy.txt",
        format="txt",
        category="test",
        access_notes="Synthetic test response only",
        intended_scope="Test",
    )
    async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
        connector = HTTPConnector([specification], client=client)
        first = await pipeline.scan([connector], staged=True)
        calls = inference.calls
        second = await pipeline.scan([connector], staged=True)
    assert first.downloaded == 1 and first.unknown == 1 and first.processed == 1
    assert second.downloaded == 0 and second.reused == 1 and second.known == 1
    assert second.skipped == 1 and inference.calls == calls and len(requests) == 2


async def test_temporary_inaccessibility_does_not_delete_history(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    await pipeline.scan([FilesystemConnector(source)])

    class Inaccessible(FilesystemConnector):
        async def fetch(self, item, previous):
            raise PermissionError("Test source denied")

    (source / "policy.txt").write_text("Users must authenticate and retain logs.")
    report = await pipeline.scan([Inaccessible(source)])
    assert report.failed == 1
    async with database() as uow:
        record = (await uow.sources.list())[0]
        assert record.status == "inaccessible" and record.revision_id is not None
        assert len(await uow.knowledge.search()) == 1
    recovered = await pipeline.scan([FilesystemConnector(source)])
    assert recovered.processed == 1 and recovered.failed == 0
    async with database() as uow:
        assert len(await uow.knowledge.search()) == 2
        assert (await uow.sources.list())[0].status == "available"


async def test_changed_processing_configuration_reuses_parsing(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    await pipeline.scan([FilesystemConnector(source)])
    pipeline.inference_fingerprint = "different-prompt-version"
    report = await pipeline.scan([FilesystemConnector(source)])
    assert report.processed == 1
    assert parser.calls == 1 and inference.calls == 2


async def test_draft_withdrawal_and_parser_replay_preserve_publication(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    path = source / "guidelines.md"

    def document(version, status):
        return (
            f"# Synthetic internal guidelines\nStatus: {status}\nVersion: {version}\n\n"
            "# Requirement\nUsers must authenticate."
        )

    path.write_text(document("2.0.0-draft.1", "Draft"))
    first = await pipeline.scan([FilesystemConnector(source)])
    assert first.processed == 1 and first.failed == 0
    async with database() as uow:
        initial_source = (await uow.sources.list())[0]
        assert initial_source.revision_id is not None
        assert initial_source.publication.publication_status == "draft"
        assert str(initial_source.publication.semantic_version) == "2.0.0-draft.1"
    unchanged = await pipeline.scan([FilesystemConnector(source)])
    assert unchanged.skipped == 1 and parser.calls == 1
    path.write_text(document("2.0.0-draft.2", "Draft"))
    assert (await pipeline.scan([FilesystemConnector(source)])).failed == 0
    path.write_text(document("2.0.0-draft.2", "Withdrawn"))
    assert (await pipeline.scan([FilesystemConnector(source)])).failed == 0
    async with database() as uow:
        current = (await uow.sources.list())[0]
        assert current.publication.publication_status == "withdrawn"
        historical = await uow.sources.revision(initial_source.revision_id)
        assert historical and historical.publication.publication_status == "draft"
        concepts = await uow.knowledge.search("authenticate")
        assert len(concepts) == 1 and len(concepts[0].versions) == 3
        assert all(
            state.governance_status == "proposed" and state.canonical_status == "not_canonical"
            for state in concepts[0].authority_states
        )
        assert all(version.semantic_version is None for version in concepts[0].versions)
    parser.version = "new-parser-with-identical-output"
    replayed = await pipeline.scan([FilesystemConnector(source)])
    assert replayed.failed == 0 and replayed.processed == 1 and parser.calls == 4
    async with database() as uow:
        concept = (await uow.knowledge.search("authenticate"))[0]
        assert len(concept.versions) == 3


async def test_publication_ties_use_the_same_status_for_source_and_governance(database, tmp_path):
    from datetime import UTC, datetime

    from sqlalchemy import update

    from ekb.adapters.persistence.postgres.models import PublicationRow
    from ekb.domain.models import (
        PublicationEvidence,
        PublicationMetadata,
        PublicationStatus,
        stable_id,
    )

    pipeline, _, _, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    assert (await pipeline.scan([FilesystemConnector(source)])).processed == 1
    async with database() as uow:
        record = (await uow.sources.list())[0]
        assert record.revision_id is not None
        revision_id = record.revision_id
        metadata = {**record.item.metadata, "draft_signal": "draft", "release_signal": "published"}
        await uow.sources.save(
            record.model_copy(
                update={"item": record.item.model_copy(update={"metadata": metadata})}
            )
        )
        low, high = sorted(("candidate-a", "candidate-b"), key=lambda f: stable_id(revision_id, f))
        for fingerprint, status, field in (
            (low, PublicationStatus.DRAFT, "draft_signal"),
            (high, PublicationStatus.PUBLISHED, "release_signal"),
        ):
            await uow.sources.save_publication(
                revision_id,
                PublicationMetadata(
                    publication_status=status,
                    evidence=(PublicationEvidence(field=field, excerpt=status.value),),
                ),
                fingerprint,
            )
        await uow.commit()
    async with database.engine.begin() as connection:
        await connection.execute(
            update(PublicationRow)
            .where(
                PublicationRow.revision_id == revision_id,
                PublicationRow.fingerprint.in_((low, high)),
            )
            .values(recorded_at=datetime(2099, 1, 1, tzinfo=UTC))
        )
    async with database() as uow:
        current = await uow.sources.get(record.item.id)
        concept = (await uow.knowledge.search("authenticate"))[0]
        snapshot = await uow.resolution.snapshot()
    assert current and current.publication.publication_status == "draft"
    assert concept.versions[0].id in snapshot.preview_version_ids


async def test_cancellation_releases_task_and_persists_scan_for_resume(database, tmp_path):
    pipeline, parser, inference, source = build(database, tmp_path)
    (source / "policy.txt").write_text("Users must authenticate.")
    entered = asyncio.Event()

    class BlockingInference(RulesInference):
        async def extract(self, segment):
            entered.set()
            await asyncio.Event().wait()
            return await super().extract(segment)

    pipeline.inference = BlockingInference()
    scan = asyncio.create_task(pipeline.scan([FilesystemConnector(source)]))
    await asyncio.wait_for(entered.wait(), 10)
    scan.cancel()
    with pytest.raises(asyncio.CancelledError):
        await scan
    async with database() as uow:
        assert (await uow.processing.tasks())[0].status == "cancelled"
        assert (await uow.processing.scans())[0].status == "cancelled"
    pipeline.inference = inference
    resumed = await pipeline.scan([FilesystemConnector(source)])
    assert resumed.processed == 1 and resumed.failed == 0
    assert parser.calls == 1 and inference.calls == 1


async def test_pdf_word_powerpoint_directory_roundtrip(database, tmp_path):
    import io

    import pymupdf
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    pipeline, parser, inference, source = build(database, tmp_path)
    statement = "Users must authenticate."
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), statement)
    pdf.save(str(source / "policy.pdf"))
    pdf.close()
    document = Document()
    document.add_paragraph(statement)
    document.save(str(source / "policy.docx"))
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = statement
    stream = io.BytesIO()
    presentation.save(stream)
    (source / "policy.pptx").write_bytes(stream.getvalue())
    report = await pipeline.scan([FilesystemConnector(source)])
    assert report.processed == 3 and report.failed == 0
    async with database() as uow:
        concepts = await uow.knowledge.search()
        assert len(concepts) == 1 and len(concepts[0].occurrences) == 3
        provenance = await uow.knowledge.provenance(concepts[0].id)
        assert {p.source.item.format for p in provenance} == {"pdf", "docx", "pptx"}
        assert all(p.revision.source_location == p.source.item.location for p in provenance)


async def test_metadata_only_sources_are_visible_without_download_or_inference(database, tmp_path):
    import httpx

    from ekb.adapters.sources import HTTPConnector
    from ekb.domain.models import SourceSpec

    pipeline, parser, inference, source = build(database, tmp_path)

    def forbidden_download(request):
        raise AssertionError("Metadata-only entries must never be downloaded")

    specification = SourceSpec(
        id="iso-metadata-test",
        publisher="ISO",
        name="Metadata only",
        location="https://example.test/standard",
        format="html",
        category="industry-standard",
        metadata_only=True,
        access_notes="No licensed text",
        intended_scope="Reference only",
    )
    async with httpx.AsyncClient(transport=httpx.MockTransport(forbidden_download)) as client:
        report = await pipeline.scan([HTTPConnector([specification], client=client)])
    assert report.skipped == 1 and report.failed == 0
    assert parser.calls == 0 and inference.calls == 0
    async with database() as uow:
        record = (await uow.sources.list())[0]
        assert record.status == "metadata_only" and record.revision_id is None
