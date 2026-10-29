import io
import json
from uuid import uuid4

import pymupdf
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from ekb.adapters.parsers import StructuredParser, split_text
from ekb.domain.models import SourceItem, SourceRevision
from ekb.domain.publication import assess_publication


async def parse(content, format, **metadata):
    source_id, revision_id = uuid4(), uuid4()
    item = SourceItem(
        id=source_id,
        connector="test",
        location="test",
        name="test",
        format=format,
        metadata=metadata,
    )
    revision = SourceRevision(
        id=revision_id, source_id=source_id, content_hash="test", artifact_key="test"
    )
    parser = StructuredParser()
    first = await parser.parse(item, revision, content)
    assert first == await parser.parse(item, revision, content)
    return first


async def test_markdown_preserves_hierarchy_and_exact_offsets():
    segments = await parse(b"# Policy\nUsers must authenticate.\n## Logs\nRetain logs.", "md")
    assert len(segments) == 2
    assert segments[1].parent_id == segments[0].id
    assert segments[1].heading == "Logs"
    assert segments[1].location["line"] == 3


async def test_pdf_text_and_selected_pages():
    document = pymupdf.open()
    for text in (
        "Users must authenticate to all public services.",
        "Operators should review security logs every quarter.",
    ):
        page = document.new_page()
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    segments = await parse(content, "pdf", pages=[2])
    assert len(segments) == 1
    assert segments[0].location["page"] == 2
    assert "should review" in segments[0].text
    assert segments[0].quality == "embedded"


async def test_docx_preserves_paragraph_table_reading_order():
    document = Document()
    document.add_heading("Policy", 1)
    document.add_paragraph("Users must authenticate.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "Retain logs"
    document.add_paragraph("Administrators should review access.")
    stream = io.BytesIO()
    document.save(stream)
    segments = await parse(stream.getvalue(), "docx")
    assert [s.text for s in segments] == [
        "Policy",
        "Users must authenticate.",
        "Control | Retain logs",
        "Administrators should review access.",
    ]
    assert segments[1].parent_id == segments[0].id
    assert segments[2].location == {"table": 2, "row": 0, "block_offset": 0}


async def test_pptx_preserves_slide_notes_and_text():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    textbox.text = "We aim to improve availability."
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "Operators must retain logs."
    stream = io.BytesIO()
    presentation.save(stream)
    segments = await parse(stream.getvalue(), "pptx")
    assert segments[0].location["slide"] == 1
    assert "improve availability" in segments[0].text
    assert segments[1].location["notes"] is True
    assert "retain logs" in segments[1].text


@pytest.mark.parametrize("format", ["pdf", "docx", "pptx", "html"])
async def test_publication_metadata_has_exact_separate_evidence(format):
    title = "Synthetic architecture standard (preview)"
    if format == "pdf":
        document = pymupdf.open()
        document.set_metadata({"title": title})
        document.new_page().insert_text((72, 72), "Users should authenticate.")
        content = document.tobytes()
        document.close()
    elif format in {"docx", "pptx"}:
        office = Document() if format == "docx" else Presentation()
        office.core_properties.title = title
        office.core_properties.version = "2.0.0-preview.2"
        stream = io.BytesIO()
        office.save(stream)
        content = stream.getvalue()
    else:
        content = (
            f"<html><head><title>{title}</title></head><body><p>Guidance.</p></body></html>"
        ).encode()
    segments = await parse(content, format)
    item = SourceItem(id=uuid4(), connector="test", location="test", name="standard", format=format)
    publication = assess_publication(item, segments)
    assert publication.publication_status == "preview"
    assert segments[0].quality == "document_metadata"
    assert publication.evidence
    for reference in publication.evidence:
        segment = next(segment for segment in segments if segment.id == reference.segment_id)
        assert segment.text[reference.start : reference.end] == reference.excerpt


async def test_oscal_selection_reaches_nested_subcategories():
    content = json.dumps(
        {
            "catalog": {
                "groups": [
                    {
                        "id": "GV",
                        "title": "Govern",
                        "controls": [
                            {
                                "id": "GV.OC",
                                "title": "Context",
                                "controls": [
                                    {
                                        "id": "GV.OC-01",
                                        "title": "Mission",
                                        "parts": [
                                            {"id": "s1", "prose": "The mission is understood."}
                                        ],
                                    },
                                    {
                                        "id": "GV.OC-02",
                                        "title": "Other",
                                        "parts": [{"id": "s2", "prose": "Not selected."}],
                                    },
                                ],
                            }
                        ],
                    }
                ]
            }
        }
    ).encode()
    segments = await parse(content, "oscal", control_ids=["GV.OC-01"])
    assert [s.text for s in segments] == ["The mission is understood."]
    assert segments[0].location["control_id"] == "GV.OC-01"
    assert segments[0].location["part_id"] == "s1"


async def test_official_xml_and_public_web_json_retain_locations():
    xml = b"<ROOT><P>Operators must maintain equipment.</P><P>Inspect annually.</P></ROOT>"
    segments = await parse(xml, "xml")
    assert len(segments) == 2 and segments[1].location["xml_element"] == 1
    web = json.dumps({"main": {"text": "<p>We aim to improve efficiency.</p>"}}).encode()
    segments = await parse(web, "web-json")
    assert segments[0].text == "We aim to improve efficiency."
    assert segments[0].location["json_path"] == "$.main.text"


def test_long_sections_split_without_losing_evidence_characters():
    text = "Users must not disable alarms. " * 100
    pieces = split_text(text, 300)
    assert len(pieces) > 1
    for offset, piece in pieces:
        assert text[offset : offset + len(piece)] == piece
        assert len(piece) <= 300


async def test_non_oscal_json_is_not_silently_treated_as_catalog():
    with pytest.raises(ValueError, match="OSCAL"):
        await parse(b'{"message":"not a catalog"}', "json")


async def test_ocr_is_opt_in_and_only_runs_on_pages_without_embedded_text(monkeypatch):
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (200, 80), "white").save(image, format="PNG")
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Users must authenticate to all public services.")
    page = document.new_page()
    page.insert_image(page.rect, stream=image.getvalue())
    content = document.tobytes()
    document.close()
    calls = []

    def ocr(page, **kwargs):
        calls.append(page.number)
        return page.get_textpage()

    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", ocr)
    source_id, revision_id = uuid4(), uuid4()
    item = SourceItem(id=source_id, connector="test", location="test", name="test", format="pdf")
    revision = SourceRevision(
        id=revision_id, source_id=source_id, content_hash="test", artifact_key="test"
    )
    default = await StructuredParser().parse(item, revision, content)
    assert calls == [] and default[-1].quality == "image_only"
    enabled = await StructuredParser(ocr=True).parse(item, revision, content)
    assert calls == [1] and enabled[-1].quality == "ocr_no_text"
